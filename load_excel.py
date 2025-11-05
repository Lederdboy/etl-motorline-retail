import pandas as pd
import os
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows
import warnings

warnings.filterwarnings('ignore', category=UserWarning, module='openpyxl')

print("="*80)
print("LOAD - Generacion de Excel con reportes y graficos")
print("="*80)

workspace = os.getcwd()

# ===== NUEVA LÓGICA: Carpeta OUTPUT + TIMESTAMP =====
output_dir = os.path.join(workspace, 'output')
if not os.path.exists(output_dir):
    os.makedirs(output_dir)
    print(f"[INFO] Carpeta 'output' creada en: {output_dir}")

timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
archivo_entrada = os.path.join(workspace, 'ventas_clean.csv')
archivo_salida = os.path.join(output_dir, f'MotorLine_Reporte_Ventas_{timestamp}.xlsx')
# ===================================================

try:
    print(f"[INFO] Leyendo CSV limpio: {archivo_entrada}")
    df = pd.read_csv(archivo_entrada, encoding='utf-8-sig')
    
    print(f"[OK] Registros leidos: {len(df)}")
    print(f"[INFO] Archivo de salida: {archivo_salida}")
    
    print(f"[INFO] Creando archivo Excel con multiples hojas...")
    
    # HOJA 1 y 2: Crear con un solo ExcelWriter
    print(f"  [1/3] Creando hoja: Datos_Ventas...")
    print(f"  [2/3] Creando hoja: Resumen_Estadisticas...")
    
    # Calcular estadisticas
    resumen = {
        'Metrica': [
            'Total Ventas',
            'Ingresos Totales',
            'Ticket Promedio',
            'Margen Promedio (%)',
            'Ventas Rentables',
            'Clientes Unicos',
            'Vendedores Activos',
            'Sucursales Activas',
            'Marcas Vendidas'
        ],
        'Valor': [
            len(df),
            f"S/. {df['Precio_venta_real'].sum():,.2f}",
            f"S/. {df['Precio_venta_real'].mean():,.2f}",
            f"{df['Porcentaje_Margen'].mean():.2f}%",
            f"{df['Venta_Rentable'].sum()} ({df['Venta_Rentable'].sum()/len(df)*100:.1f}%)",
            df['ID_Cliente'].nunique(),
            df['ID_Empleado'].nunique(),
            df['ID_Sucursal'].nunique(),
            df['Marca'].nunique()
        ]
    }
    df_resumen = pd.DataFrame(resumen)
    
    # Top 5 Clientes
    top_clientes = df.groupby('Cliente_NombreCompleto')['Precio_venta_real'].agg(['sum', 'count']).reset_index()
    top_clientes.columns = ['Cliente', 'Total_Comprado', 'Cantidad_Compras']
    top_clientes = top_clientes.sort_values('Total_Comprado', ascending=False).head(5)
    top_clientes['Total_Comprado'] = top_clientes['Total_Comprado'].apply(lambda x: f"S/. {x:,.2f}")
    
    # Top 5 Vendedores
    top_vendedores = df.groupby('Vendedor_NombreCompleto')['Precio_venta_real'].agg(['sum', 'count']).reset_index()
    top_vendedores.columns = ['Vendedor', 'Total_Vendido', 'Cantidad_Ventas']
    top_vendedores = top_vendedores.sort_values('Total_Vendido', ascending=False).head(5)
    top_vendedores['Total_Vendido'] = top_vendedores['Total_Vendido'].apply(lambda x: f"S/. {x:,.2f}")
    
    # Ventas por Sucursal
    ventas_sucursal = df.groupby('Nombre_Sucursal')['Precio_venta_real'].agg(['sum', 'count']).reset_index()
    ventas_sucursal.columns = ['Sucursal', 'Total_Ventas', 'Cantidad']
    ventas_sucursal = ventas_sucursal.sort_values('Total_Ventas', ascending=False)
    ventas_sucursal['Total_Ventas'] = ventas_sucursal['Total_Ventas'].apply(lambda x: f"S/. {x:,.2f}")
    
    # Ventas por Marca
    ventas_marca = df.groupby('Marca')['Precio_venta_real'].agg(['sum', 'count']).reset_index()
    ventas_marca.columns = ['Marca', 'Total_Ventas', 'Cantidad']
    ventas_marca = ventas_marca.sort_values('Total_Ventas', ascending=False)
    ventas_marca['Total_Ventas'] = ventas_marca['Total_Ventas'].apply(lambda x: f"S/. {x:,.2f}")
    
    # Ventas Mensuales
    ventas_mensuales = df.groupby(['Anio_Venta', 'Mes_Venta', 'Mes_Nombre'])['Precio_venta_real'].agg(['sum', 'count']).reset_index()
    ventas_mensuales.columns = ['Anio', 'Mes_Num', 'Mes', 'Total_Ventas', 'Cantidad']
    ventas_mensuales = ventas_mensuales.sort_values(['Anio', 'Mes_Num'])
    ventas_mensuales['Total_Ventas'] = ventas_mensuales['Total_Ventas'].apply(lambda x: f"S/. {x:,.2f}")
    
    # Escribir HOJA 1 y HOJA 2 con un solo ExcelWriter
    with pd.ExcelWriter(archivo_salida, engine='openpyxl') as writer:
        # Hoja 1: Datos completos
        df.to_excel(writer, sheet_name='Datos_Ventas', index=False)
        
        # Hoja 2: Resumen con workbook actual
        workbook = writer.book
        worksheet = workbook.create_sheet('Resumen_Estadisticas')
        
        # Escribir datos manualmente en la hoja
        current_row = 1
        
        # Resumen general
        for r_idx, row in enumerate(dataframe_to_rows(df_resumen, index=False, header=True), 1):
            for c_idx, value in enumerate(row, 1):
                worksheet.cell(row=current_row, column=c_idx, value=value)
            current_row += 1
        
        current_row += 2  # Espaciado
        
        # Top Clientes
        for r_idx, row in enumerate(dataframe_to_rows(top_clientes, index=False, header=True), 1):
            for c_idx, value in enumerate(row, 1):
                worksheet.cell(row=current_row, column=c_idx, value=value)
            current_row += 1
        
        current_row += 2
        
        # Top Vendedores
        for r_idx, row in enumerate(dataframe_to_rows(top_vendedores, index=False, header=True), 1):
            for c_idx, value in enumerate(row, 1):
                worksheet.cell(row=current_row, column=c_idx, value=value)
            current_row += 1
        
        current_row += 2
        
        # Ventas por Sucursal
        for r_idx, row in enumerate(dataframe_to_rows(ventas_sucursal, index=False, header=True), 1):
            for c_idx, value in enumerate(row, 1):
                worksheet.cell(row=current_row, column=c_idx, value=value)
            current_row += 1
        
        current_row += 2
        
        # Ventas por Marca
        for r_idx, row in enumerate(dataframe_to_rows(ventas_marca, index=False, header=True), 1):
            for c_idx, value in enumerate(row, 1):
                worksheet.cell(row=current_row, column=c_idx, value=value)
            current_row += 1
        
        current_row += 2
        
        # Ventas Mensuales
        for r_idx, row in enumerate(dataframe_to_rows(ventas_mensuales, index=False, header=True), 1):
            for c_idx, value in enumerate(row, 1):
                worksheet.cell(row=current_row, column=c_idx, value=value)
            current_row += 1
    
    # HOJA 3: DASHBOARD CON GRAFICOS
    print(f"  [3/3] Creando hoja: Dashboard_Graficos...")
    
    # Preparar datos para graficos
    ventas_mes_graficos = df.groupby(['Anio_Venta', 'Mes_Venta'])['Precio_venta_real'].sum().reset_index()
    ventas_mes_graficos['Periodo'] = ventas_mes_graficos['Anio_Venta'].astype(str) + '-' + ventas_mes_graficos['Mes_Venta'].astype(str).str.zfill(2)
    
    ventas_sucursal_grafico = df.groupby('Nombre_Sucursal')['Precio_venta_real'].sum().reset_index()
    ventas_sucursal_grafico = ventas_sucursal_grafico.sort_values('Precio_venta_real', ascending=False)
    
    top_modelos = df.groupby('Vehiculo_Completo')['Precio_venta_real'].count().reset_index()
    top_modelos.columns = ['Modelo', 'Cantidad']
    top_modelos = top_modelos.sort_values('Cantidad', ascending=False).head(10)
    
    metodos_pago = df.groupby('Metodo_Pago')['Precio_venta_real'].sum().reset_index()
    
    # Cargar el workbook para agregar la hoja 3
    wb = load_workbook(archivo_salida)
    ws3 = wb.create_sheet('Dashboard_Graficos')
    
    # Titulo
    ws3['B1'] = 'DASHBOARD DE VENTAS - MOTORLINE S.A.C.'
    ws3['B1'].font = Font(bold=True, size=16, color="FFFFFF")
    ws3['B1'].fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    ws3['B1'].alignment = Alignment(horizontal="center", vertical="center")
    ws3.merge_cells('B1:O1')
    
    # Escribir datos para graficos
    # Ventas mensuales
    row = 2
    for r_idx, data_row in enumerate(dataframe_to_rows(ventas_mes_graficos[['Periodo', 'Precio_venta_real']], index=False, header=True), row):
        for c_idx, value in enumerate(data_row, 2):
            ws3.cell(row=r_idx, column=c_idx, value=value)
    
    # Ventas por sucursal
    row = 2
    for r_idx, data_row in enumerate(dataframe_to_rows(ventas_sucursal_grafico, index=False, header=True), row):
        for c_idx, value in enumerate(data_row, 6):
            ws3.cell(row=r_idx, column=c_idx, value=value)
    
    # Top modelos
    row = 2
    for r_idx, data_row in enumerate(dataframe_to_rows(top_modelos, index=False, header=True), row):
        for c_idx, value in enumerate(data_row, 10):
            ws3.cell(row=r_idx, column=c_idx, value=value)
    
    # Metodos de pago
    row = 2
    for r_idx, data_row in enumerate(dataframe_to_rows(metodos_pago, index=False, header=True), row):
        for c_idx, value in enumerate(data_row, 14):
            ws3.cell(row=r_idx, column=c_idx, value=value)
    
    # FORMATEAR HOJAS
    print(f"[INFO] Aplicando formato y graficos...")
    
    # Formatear Hoja 1: Datos_Ventas
    ws1 = wb['Datos_Ventas']
    for cell in ws1[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Formatear Hoja 2: Resumen_Estadisticas (encabezados en cada tabla)
    ws2 = wb['Resumen_Estadisticas']
    for row in ws2.iter_rows(min_row=1, max_row=ws2.max_row):
        if row[0].value and str(row[0].value) in ['Metrica', 'Cliente', 'Vendedor', 'Sucursal', 'Marca', 'Anio']:
            for cell in row:
                if cell.value:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
                    cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Formatear encabezados de Hoja 3
    for cell in ws3[2]:
        if cell.value:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Grafico 1: Ventas Mensuales (Linea)
    chart1 = LineChart()
    chart1.title = "Ventas Mensuales"
    chart1.y_axis.title = "Monto (S/.)"
    chart1.x_axis.title = "Periodo"
    
    data1 = Reference(ws3, min_col=3, min_row=2, max_row=2+len(ventas_mes_graficos))
    cats1 = Reference(ws3, min_col=2, min_row=3, max_row=2+len(ventas_mes_graficos))
    chart1.add_data(data1, titles_from_data=True)
    chart1.set_categories(cats1)
    chart1.width = 15
    chart1.height = 10
    ws3.add_chart(chart1, "B20")
    
    # Grafico 2: Ventas por Sucursal (Barras)
    chart2 = BarChart()
    chart2.title = "Ventas por Sucursal"
    chart2.y_axis.title = "Monto (S/.)"
    chart2.x_axis.title = "Sucursal"
    
    data2 = Reference(ws3, min_col=7, min_row=2, max_row=2+len(ventas_sucursal_grafico))
    cats2 = Reference(ws3, min_col=6, min_row=3, max_row=2+len(ventas_sucursal_grafico))
    chart2.add_data(data2, titles_from_data=True)
    chart2.set_categories(cats2)
    chart2.width = 15
    chart2.height = 10
    ws3.add_chart(chart2, "J20")
    
    # Grafico 3: Top 10 Modelos (Barras horizontales)
    chart3 = BarChart()
    chart3.type = "bar"
    chart3.title = "Top 10 Modelos Vendidos"
    chart3.y_axis.title = "Modelo"
    chart3.x_axis.title = "Cantidad"
    
    data3 = Reference(ws3, min_col=11, min_row=2, max_row=2+len(top_modelos))
    cats3 = Reference(ws3, min_col=10, min_row=3, max_row=2+len(top_modelos))
    chart3.add_data(data3, titles_from_data=True)
    chart3.set_categories(cats3)
    chart3.width = 15
    chart3.height = 12
    ws3.add_chart(chart3, "B40")
    
    # Grafico 4: Metodos de Pago (Torta)
    chart4 = PieChart()
    chart4.title = "Distribucion Metodos de Pago"
    
    data4 = Reference(ws3, min_col=15, min_row=2, max_row=2+len(metodos_pago))
    cats4 = Reference(ws3, min_col=14, min_row=3, max_row=2+len(metodos_pago))
    chart4.add_data(data4, titles_from_data=True)
    chart4.set_categories(cats4)
    chart4.width = 12
    chart4.height = 10
    ws3.add_chart(chart4, "J40")
    
    # Guardar Excel
    wb.save(archivo_salida)
    
    print(f"")
    print(f"[OK] Excel generado exitosamente!")
    print(f"[OK] Ubicacion: {archivo_salida}")
    print(f"")
    print(f"CONTENIDO DEL ARCHIVO:")
    print(f"  Hoja 1: Datos_Ventas ({len(df)} registros, {len(df.columns)} columnas)")
    print(f"  Hoja 2: Resumen_Estadisticas (9 metricas, 5 tablas)")
    print(f"  Hoja 3: Dashboard_Graficos (4 graficos)")
    print(f"")
    print(f"GRAFICOS INCLUIDOS:")
    print(f"  1. Ventas Mensuales (Grafico de Linea)")
    print(f"  2. Ventas por Sucursal (Grafico de Barras)")
    print(f"  3. Top 10 Modelos Vendidos (Barras Horizontales)")
    print(f"  4. Distribucion Metodos de Pago (Grafico Circular)")
    print("="*80)
    
except FileNotFoundError:
    print(f"[ERROR] No se encuentra el archivo {archivo_entrada}")
    print(f"[INFO] Asegurate de ejecutar transform_data.py primero")
    exit(1)
    
except Exception as e:
    print(f"[ERROR] Error al generar Excel: {e}")
    import traceback
    traceback.print_exc()
    exit(1)